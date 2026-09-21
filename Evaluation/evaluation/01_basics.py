# -*- coding: utf-8 -*-
"""
Auto-extracted from evaluation.ipynb
Section 1/5: 评估函数
Slug: basics
Note: outputs and cell magic were dropped; source code only.
"""


# === code block 1/4 ===
# from rouge import rouge 
import jieba
from nltk.translate.bleu_score import sentence_bleu
import json
from rouge_chinese import Rouge
import jieba
import pandas as pd
import cntext as ct
from transformers import T5Tokenizer, T5ForConditionalGeneration,AutoTokenizer,AutoModelForCausalLM,AutoModel
from gensim.models import KeyedVectors,word2vec
from sklearn.metrics.pairwise import cosine_similarity
import torch
import openpyxl
rouge = Rouge()

# === code block 2/4 ===
model = KeyedVectors.load_word2vec_format('../Medical.txt', binary=False) #词向量嵌入模型加载

# === code block 3/4 ===
#blue值-均值
def cumulative_bleu(total,reference, candidate):
    bleu_1_gram = sentence_bleu(reference, candidate, weights=(1, 0, 0, 0))
    bleu_2_gram = sentence_bleu(reference, candidate, weights=(0.5, 0.5, 0, 0))
    bleu_3_gram = sentence_bleu(reference, candidate, weights=(0.33, 0.33, 0.33, 0))
    bleu_4_gram = sentence_bleu(reference, candidate, weights=(0.25, 0.25, 0.25, 0.25))
    total[0] = total[0] + bleu_1_gram
    total[1] = total[1] + bleu_2_gram
    total[2] = total[2] + bleu_3_gram
    total[3] = total[3] + bleu_4_gram
    return total

#rouge值
def roug(data_rouge):
    hyps, refs = map(list, zip(*[[' '.join(jieba.cut(d['hyp'])), ' '.join(jieba.cut(d['ref']))] for d in data_rouge]))
    scores = rouge.get_scores(hyps, refs, avg=True) #rouge值
    rouge_1 = scores['rouge-1']['r']
    rouge_2 = scores['rouge-2']['r']
    rouge_l = scores['rouge-l']['r']
    rouge_score = [rouge_1,rouge_2,rouge_l]
    return rouge_score

#文本的可读性
def keduc(data_rouge,m):
    score = 0
    for d in data_rouge:
        result = ct.readability(d['hyp'], lang='chinese')
        result = sum(result.values()) / len(result)
        # print(result)
        # result = result['readability3']
        score = score + result
    score = score/m
    return score

def zhuijia(values, name):
    file_full_path = f'./{name}.xlsx'
    sheet_name = 'Sheet1'
    wb = openpyxl.load_workbook(file_full_path)
    ws = wb[sheet_name]
    max_row_num = ws.max_row
    max_col_num = ws.max_column
    ws._current_row = max_row_num
    ws.append(values)
    wb.save(file_full_path)

def sim(data_rouge,s):
    score = 0
    for i in data_rouge:
        m = i['hyp']
        c = i['ref']
        m_words_vec = []
        c_words_vec = []
        #分词，取向量相加
        m = list(jieba.cut(m))
        for i,words in enumerate(m):
            try:
                if i==0:
                    temp = model[words]
                else:
                    temp = temp+model[words]
            except KeyError:
                s = s-1
                continue
        c = list(jieba.cut(c))
        for i,words in enumerate(c):
            try:
                if i==0:
                    temp1 = model[words]
                else:
                    temp1 = temp1+model[words]
            except KeyError:
                s = s-1
                continue
        #取平均
        temp = temp/len(m)
        temp = list(temp)
        temp1 = temp1/len(c)
        temp1 = list(temp1)
        all_vec = [temp,temp1]
        #计算
        cos_similarity = cosine_similarity(all_vec)
        cos_similarity = cos_similarity[0][1]
        score = score + cos_similarity
    score = score/s
    return score


def much_bleu_rouge(file, name):
    rouge = Rouge()
    with open(file,'r') as f:
        data = f.readlines()
    dialogues = []
    h = 0
    for i in data:
        i = json.loads(i)
        dialogues = i['conversation']
        m = len(dialogues)
        if m==0:
            continue
        total = [0,0,0,0]
        data_rouge = []
        dict_rouge = {}
        for j in dialogues:
            candidate = qianwen(j['human']) #模型结果
            reference = []
            reference.append(list(jieba.cut(j['assistant']))) #真实答案
            total = cumulative_bleu(total,reference,candidate) #blue值
            dict_rouge['hyp'] = candidate
            dict_rouge['ref'] = j['assistant']
            data_rouge.append(dict_rouge)
        bleu_score = [mm/m for mm in total] #blue值-均值
        rouge_score = roug(data_rouge) #rouge值
        readability = keduc(data_rouge,m)
        similarity = sim(data_rouge,m)
        score = bleu_score + rouge_score
        score.append(readability)
        score.append(similarity)
        score.insert(0, h)
        #追加插入
        zhuijia(score, name)
        h = h+1
        print(h)

# === code block 4/4 ===
def sim(data_rouge,s):
    score = 0
    for i in data_rouge:
        m = i['hyp']
        c = i['ref']
        m_words_vec = []
        c_words_vec = []
        #分词，取向量相加
        m = list(jieba.cut(m))
        for i,words in enumerate(m):
            try:
                if i==0:
                    temp = model[words]
                else:
                    temp = temp+model[words]
            except KeyError:
                s = s-1
                continue
        c = list(jieba.cut(c))
        for i,words in enumerate(c):
            try:
                if i==0:
                    temp1 = model[words]
                else:
                    temp1 = temp1+model[words]
            except KeyError:
                s = s-1
                continue
        #取平均
        temp = temp/len(m)
        temp = list(temp)
        temp1 = temp1/len(c)
        temp1 = list(temp1)
        all_vec = [temp,temp1]
        #计算
        cos_similarity = cosine_similarity(all_vec)
        cos_similarity = cos_similarity[0][1]
        score = score + cos_similarity
    score = score/s
    return score


def much_bleu_rouge(file, name):
    rouge = Rouge()
    with open(file,'r') as f:
        data = f.readlines()
    dialogues = []
    h = 0
    for i in data:
        i = json.loads(i)
        dialogues = i['conversation']
        m = len(dialogues)
        if m==0:
            continue
        total = [0,0,0,0]
        data_rouge = []
        dict_rouge = {}
        for j in dialogues:
            candidate = qianwen(j['human']) #模型结果
            dict_rouge['hyp'] = candidate
            dict_rouge['ref'] = j['assistant']
            data_rouge.append(dict_rouge)

        similarity = sim(data_rouge,m)
        print(data_rouge)
        print(similarity)
